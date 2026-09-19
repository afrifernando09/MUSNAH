import streamlit as st
import pandas as pd
import google.generativeai as genai
import json
import re
from PIL import Image
import io
import openpyxl
from openpyxl.styles import PatternFill

st.set_page_config(page_title="Rekap Catatan Stok", layout="wide")

st.title("📝 Aplikasi Rekap Catatan Stok Otomatis")
st.write("Upload banyak gambar catatan sekaligus dan file MASTER STOK Anda di sini.")

with st.sidebar:
    st.header("Pengaturan")
    api_key = st.text_input("Masukkan Gemini API Key", type="password")
    st.markdown("[Dapatkan API Key Gratis di sini](https://aistudio.google.com/app/apikey)")

col1, col2 = st.columns(2)
with col1:
    master_file = st.file_uploader("1. Upload MASTER STOK (.xls / .xlsx)", type=["xls", "xlsx"])
with col2:
    image_files = st.file_uploader("2. Upload Gambar Catatan (.jpg / .png)", type=["jpg", "jpeg", "png"], accept_multiple_files=True)

def get_smarter_match(plu_tulis, deskripsi_tulis, is_pck, df_master):
    plu_bersih = str(plu_tulis).replace('D', '0').replace('O', '0').strip()
    desk_tulis = str(deskripsi_tulis).upper().strip()
    pck_units = ['PCS', 'PCK', 'IKT', 'CTN', 'CUP']
    
    def evaluate(df_sub):
        if df_sub.empty: return None
        df_sub = df_sub.copy()
        df_sub['Has_Stock'] = (df_sub['STOK'] != 0).astype(int)
        df_sub['Unit_Match'] = df_sub['UNIT'].isin(pck_units).astype(int) if is_pck else 1
        return df_sub.sort_values(by=['Has_Stock', 'Unit_Match'], ascending=[False, False]).iloc[0]

    exact_plu = df_master[df_master['PLU'] == plu_bersih]
    if not exact_plu.empty: return evaluate(exact_plu)
        
    if len(plu_bersih) >= 5:
        sim_plu = df_master[df_master['PLU'].str.contains(plu_bersih, na=False)]
        if not sim_plu.empty: return evaluate(sim_plu)
            
    if desk_tulis.startswith('B/O') or desk_tulis.startswith('BO '):
        df_bo = df_master[df_master['DESKRIPSI'].str.contains('BUAH OLAHAN', na=False, case=False)]
        if 'MIX' in desk_tulis:
            df_bo_mix = df_bo[df_bo['DESKRIPSI'].str.contains('CAMPUR', na=False, case=False)]
            if not df_bo_mix.empty: return evaluate(df_bo_mix)
        else:
            sisa_teks = re.sub(r'^(B/O|BO)\s*', '', desk_tulis)
            keywords = [k for k in re.findall(r'[A-Z0-9]+', sisa_teks) if len(k) > 2]
            if keywords:
                df_bo_copy = df_bo.copy()
                df_bo_copy['score'] = df_bo_copy['DESKRIPSI'].apply(lambda x: sum(1 for k in keywords if k in str(x).upper()))
                max_score = df_bo_copy['score'].max()
                if max_score > 0: return evaluate(df_bo_copy[df_bo_copy['score'] == max_score])
    else:
        replace_dict = {'B BOMBAI': 'BAWANG BOMBAY', 'B MERAH': 'BAWANG MERAH', 'B PUTIH': 'BAWANG PUTIH', 'B BIMA': 'BAWANG BIRMA'}
        for k, v in replace_dict.items():
            if desk_tulis.startswith(k) or k in desk_tulis:
                desk_tulis = desk_tulis.replace(k, v)
                
        keywords = [k for k in re.findall(r'[A-Z0-9]+', desk_tulis) if len(k) > 2]
        if keywords:
            df_master_copy = df_master.copy()
            df_master_copy['score'] = df_master_copy['DESKRIPSI'].apply(lambda x: sum(1 for k in keywords if k in str(x).upper()))
            df_master_copy.loc[df_master_copy['DESKRIPSI'].str.contains('BUAH OLAHAN', na=False, case=False), 'score'] -= 1
            max_score = df_master_copy['score'].max()
            if max_score > 0: return evaluate(df_master_copy[df_master_copy['score'] == max_score])
            
    return None

def apply_grouping(df_results, max_sum=495000):
    if df_results.empty: return df_results
    df_results['Kelompok'] = 0
    current_group = 1
    current_sum = 0
    for index, row in df_results.iterrows():
        if current_sum + row['Total Rupiah'] > max_sum and current_sum > 0:
            current_group += 1
            current_sum = row['Total Rupiah']
        else:
            current_sum += row['Total Rupiah']
        df_results.at[index, 'Kelompok'] = current_group
    return df_results

# Fungsi untuk Web Streamlit
def color_groups(row):
    colors = ['#e6f2ff', '#e6ffe6', '#ffffe6', '#ffe6e6', '#f9e6ff', '#e6ffff', '#fff0e6']
    color = colors[(row['Kelompok'] - 1) % len(colors)]
    return [f'background-color: {color}'] * len(row)

def format_indo(val):
    return f"{val:,.0f}".replace(',', '.')
    
def format_rupiah(val):
    return f"Rp {val:,.0f}".replace(',', '.')

if st.button("🚀 Proses Semua Gambar Sekarang", use_container_width=True):
    if not api_key or not master_file or not image_files:
        st.error("⚠️ Pastikan API Key, File Excel, dan setidaknya satu Gambar telah diunggah!")
    else:
        with st.spinner(f'Memproses {len(image_files)} gambar...'):
            try:
                genai.configure(api_key=api_key)
                model = genai.GenerativeModel('gemini-3.5-flash')
                
                df_master = pd.read_excel(master_file, sheet_name=0)
                df_master['PLU'] = df_master['PLU'].astype(str).str.strip()
                df_master['STOK'] = pd.to_numeric(df_master.get('STOK', 0), errors='coerce').fillna(0)
                df_master['ACOST'] = pd.to_numeric(df_master.get('ACOST', 0), errors='coerce').fillna(0)
                df_master['FRAC'] = pd.to_numeric(df_master.get('FRAC', 1), errors='coerce').fillna(1).replace(0, 1)
                
                results = []
                
                for img_file in image_files:
                    img = Image.open(img_file)
                    prompt = """
                    Ekstrak data dari gambar. Output JSON array of objects:
                    [
                      {
                        "plu_tulis": "KODE PLU (angka di kolom pertama)",
                        "deskripsi": "Deskripsi barang",
                        "qty_expr": "Operasi matematika di kolom terakhir (jangan dihitung, tulis aslinya misal '4.722+0.380+2126')"
                      }
                    ]
                    Hanya berikan JSON murni.
                    """
                    response = model.generate_content([img, prompt])
                    extracted_data = json.loads(response.text.replace("```json", "").replace("```", "").strip())
                    
                    for item in extracted_data:
                        plu = str(item.get('plu_tulis', '')).strip()
                        desk = str(item.get('deskripsi', '')).strip()
                        qty_asli = str(item.get('qty_expr', '')).strip()
                        
                        # 1. Hilangkan titik pada QTY
                        qty_clean = qty_asli.replace('.', '')
                        is_pck = bool(re.search(r'\b(pcs|pck)\b', qty_clean.lower()) or re.search(r'\b(pcs|pck)\b', desk.lower()))
                        
                        # 2. Hitung dengan mengabaikan angka 0 di depan (mengubah 0380 jadi 380)
                        try:
                            clean_expr = re.sub(r'[^\d+\-*/ ]', '', qty_clean)
                            clean_expr = re.sub(r'\b0+(\d)', r'\1', clean_expr) # Hapus 0 di depan angka
                            total_qty = eval(clean_expr) if clean_expr.strip() else 0
                        except:
                            total_qty = 0
                        
                        match = get_smarter_match(plu, desk, is_pck, df_master)
                        
                        if match is not None:
                            acost = float(match['ACOST'])
                            frac = float(match['FRAC'])
                            total_rupiah = (total_qty * acost) / frac
                            
                            results.append({
                                'Kelompok': 0,
                                'PLU Tulis': plu,
                                'Deskripsi Catatan': desk,
                                'Qty Catatan': qty_asli,
                                'Total Qty': total_qty, # KOLOM BARU TOTAL QTY
                                'Deskripsi Master': match['DESKRIPSI'],
                                'Unit': match['UNIT'],
                                'Stok System': match['STOK'],
                                'Total Rupiah': round(total_rupiah, 2),
                                'Sumber Gambar': img_file.name
                            })
                        else:
                            results.append({
                                'Kelompok': 0, 'PLU Tulis': plu, 'Deskripsi Catatan': desk, 
                                'Qty Catatan': qty_asli, 'Total Qty': total_qty,
                                'Deskripsi Master': 'TIDAK DITEMUKAN', 
                                'Unit': '-', 'Stok System': 0, 'Total Rupiah': 0.0,
                                'Sumber Gambar': img_file.name
                            })
                
                df_result = pd.DataFrame(results)
                df_result = apply_grouping(df_result, max_sum=495000)
                
                st.success("Berhasil memproses seluruh gambar!")
                
                # Tampilkan di Web (dengan format titik)
                st.dataframe(df_result.style.apply(color_groups, axis=1).format({
                    'Total Qty': format_indo, 
                    'Stok System': format_indo,
                    'Total Rupiah': format_rupiah
                }), use_container_width=True)
                
                # EXCEL EXPORT (DENGAN WARNA DAN FORMAT ANGKA)
                output = io.BytesIO()
                with pd.ExcelWriter(output, engine='openpyxl') as writer:
                    df_result.to_excel(writer, index=False, sheet_name='Rekap')
                    
                    workbook = writer.book
                    worksheet = writer.sheets['Rekap']
                    
                    # Daftar kode warna HEX untuk Excel
                    hex_colors = ['E6F2FF', 'E6FFE6', 'FFFFE6', 'FFE6E6', 'F9E6FF', 'E6FFFF', 'FFF0E6']
                    fills = [PatternFill(start_color=c, end_color=c, fill_type='solid') for c in hex_colors]
                    
                    # Memproses setiap baris untuk mewarnai dan memberi titik pemisah
                    for row_idx, row in enumerate(df_result.itertuples(), start=2): 
                        group_num = row.Kelompok
                        fill = fills[(group_num - 1) % len(fills)]
                        
                        for col_idx in range(1, len(df_result.columns) + 1):
                            cell = worksheet.cell(row=row_idx, column=col_idx)
                            cell.fill = fill
                            
                            # Format angka agar bisa dijumlahkan di Excel namun tetap memiliki pemisah
                            col_name = df_result.columns[col_idx - 1]
                            if col_name in ['Total Qty', 'Stok System']:
                                cell.number_format = '#,##0' # Format Ribuan Excel
                            elif col_name == 'Total Rupiah':
                                cell.number_format = '"Rp" #,##0'
                                
                output.seek(0)
                
                st.download_button(
                    label="📥 Download Hasil Excel (Berwarna)",
                    data=output,
                    file_name="Hasil_Rekap_Dikelompokkan_V3.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
            except Exception as e:
                st.error(f"Terjadi kesalahan: {e}")
